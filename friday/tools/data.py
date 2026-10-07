"""Spreadsheet, finance-analysis and knowledge-ingest tools."""
from __future__ import annotations

import csv
import math
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Literal

from pydantic import BaseModel, Field

from friday import finance
from friday.security.policy import Classification, Effect, Risk
from friday.tools.base import Tool, ToolContext, ToolResult, Verification
from friday.tools.pathpolicy import PathDenied


def read_table(path: Path, sheet: str | None = None, max_rows: int = 200_000) -> tuple[list[str], list[list], list[str]]:
    """-> (header, rows, sheet_names). Supports csv/tsv/xlsx/xlsm."""
    sfx = path.suffix.lower()
    if sfx in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        names = wb.sheetnames
        ws = wb[sheet] if sheet else wb.active
        it = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(it, [])]
        rows = []
        for i, r in enumerate(it):
            if i >= max_rows:
                break
            if any(c is not None and c != "" for c in r):
                rows.append(list(r))
        return header, rows, names
    if sfx in (".csv", ".tsv", ".txt"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            sample = f.read(4096); f.seek(0)
            delim = "\t" if sfx == ".tsv" else (csv.Sniffer().sniff(sample, ",;\t|").delimiter if sample else ",")
            rd = csv.reader(f, delimiter=delim)
            header = [h.strip() for h in next(rd, [])]
            rows = []
            for i, r in enumerate(rd):
                if i >= max_rows:
                    break
                if any(c.strip() for c in r):
                    rows.append(r)
        return header, rows, []
    raise ValueError(f"unsupported file type '{sfx}' (csv, tsv, xlsx supported)")


def to_number(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip()
    neg = s.startswith("(") and s.endswith(")")
    s = re.sub(r"[^\d.\-]", "", s.replace(",", ""))
    if s in ("", "-", "."):
        return None
    try:
        x = float(s)
    except ValueError:
        return None
    return -x if neg else x


def to_date(v) -> datetime | None:
    if isinstance(v, datetime):
        return v
    return finance._dt(v) if v else None


def _col(header: list[str], name: str) -> int:
    low = [h.lower() for h in header]
    if name.lower() in low:
        return low.index(name.lower())
    near = [i for i, h in enumerate(low) if name.lower() in h]
    if len(near) == 1:
        return near[0]
    raise ValueError(f"column '{name}' not found. Columns: {', '.join(header)}")


class _Sheet(Tool):
    group = "files"

    def _path(self, ctx, raw):
        return ctx.svc.paths.check_read(raw)

    def classify(self, args, ctx):
        try:
            self._path(ctx, args.path)
        except PathDenied as e:
            return Classification(blocked=str(e))
        return Classification()


class SpreadsheetSummary(_Sheet):
    name = "spreadsheet_summary"
    scope = "filesystem.read"
    description = "Describe a CSV/XLSX file: sheets, columns, row count, numeric column totals, and a few sample rows."

    class Args(BaseModel):
        path: str
        sheet: str | None = None

    def run(self, args, ctx):
        p = self._path(ctx, args.path)
        if not p.is_file():
            return ToolResult.fail(f"{p} is not a file")
        header, rows, sheets = read_table(p, args.sheet)
        numeric = {}
        for i, h in enumerate(header):
            vals = [to_number(r[i]) for r in rows if i < len(r)]
            nums = [v for v in vals if v is not None]
            if h and len(nums) >= max(1, 0.7 * len([v for v in (r[i] if i < len(r) else None for r in rows) if v not in (None, "")])) and nums:
                numeric[h] = {"count": len(nums), "sum": round(sum(nums), 4), "min": min(nums), "max": max(nums)}
        sample = [dict(zip(header, r)) for r in rows[:3]]
        return ToolResult.success(f"{p.name}: {len(rows)} data rows, {len(header)} columns" + (f", sheets: {', '.join(sheets)}" if sheets else ""),
                                  path=str(p), columns=header, rows=len(rows), sheets=sheets, numeric_columns=numeric,
                                  sample=sample)


class SpreadsheetCalc(_Sheet):
    name = "spreadsheet_calc"
    scope = "filesystem.read"
    description = ("Calculate sum/mean/min/max/count of a column in a CSV/XLSX, optionally filtered to a month "
                   "(YYYY-MM via date_column), filtered by another column's value, or grouped by a column.")

    class Args(BaseModel):
        path: str
        column: str = Field(description="Column to aggregate (e.g. 'Revenue')")
        op: Literal["sum", "mean", "min", "max", "count"] = "sum"
        sheet: str | None = None
        date_column: str | None = None
        month: str | None = Field(default=None, description="YYYY-MM; requires date_column")
        filter_column: str | None = None
        filter_equals: str | None = None
        group_by: str | None = None

    def run(self, args, ctx):
        p = self._path(ctx, args.path)
        header, rows, _ = read_table(p, args.sheet)
        ci = _col(header, args.column)
        sel = rows
        notes = []
        if args.month:
            if not args.date_column:
                return ToolResult.fail("month requires date_column")
            di = _col(header, args.date_column)
            y, m = (int(x) for x in args.month.split("-"))
            keep, undated = [], 0
            for r in sel:
                d = to_date(r[di]) if di < len(r) else None
                if d is None:
                    undated += 1
                elif (d.year, d.month) == (y, m):
                    keep.append(r)
            if undated:
                notes.append(f"{undated} row(s) had no readable date and were excluded")
            sel = keep
        if args.filter_column:
            fi = _col(header, args.filter_column)
            sel = [r for r in sel if fi < len(r) and str(r[fi]).strip().lower() == str(args.filter_equals or "").strip().lower()]

        def agg(rs):
            vals = [to_number(r[ci]) for r in rs if ci < len(r)]
            nums = [v for v in vals if v is not None]
            if args.op == "count":
                return len(nums)
            if not nums:
                return None
            return {"sum": sum, "mean": lambda x: sum(x) / len(x), "min": min, "max": max}[args.op](nums)

        skipped = sum(1 for r in sel if ci >= len(r) or to_number(r[ci]) is None)
        if skipped:
            notes.append(f"{skipped} row(s) had no numeric value in '{header[ci]}'")
        if args.group_by:
            gi = _col(header, args.group_by)
            groups: dict[str, list] = defaultdict(list)
            for r in sel:
                groups[str(r[gi]) if gi < len(r) else ""].append(r)
            res = {k: agg(v) for k, v in sorted(groups.items())}
            summ = f"{args.op} of {header[ci]} by {header[gi]}: " + ", ".join(f"{k}={v:,.2f}" if isinstance(v, float) else f"{k}={v}" for k, v in res.items())
            return ToolResult.success(summ, groups=res, rows_used=len(sel), notes=notes, path=str(p))
        v = agg(sel)
        txt = "no numeric data matched" if v is None else (f"{v:,.2f}" if isinstance(v, float) else str(v))
        return ToolResult.success(f"{args.op} of {header[ci]} = {txt} ({len(sel)} rows used){'; ' + '; '.join(notes) if notes else ''}",
                                  value=v, rows_used=len(sel), notes=notes, path=str(p), column=header[ci])


class FinanceJournal(Tool):
    name = "finance_journal_analyze"
    scope = "finance.analysis"
    group = "finance"
    description = ("Analyse a trading journal (CSV/XLSX: needs a profit column; optional date, symbol, side, size, risk, "
                   "setup, timeframe, session). Reports win rate, expectancy, profit factor, drawdown, breakdowns by "
                   "session/weekday/setup/timeframe and behavioural patterns, separated into observations, "
                   "interpretations and recommendations. Analysis only; not advice.")

    class Args(BaseModel):
        path: str
        last: int | None = Field(default=None, ge=1, description="Only the most recent N trades")
        symbol: str | None = Field(default=None, description="e.g. XAUUSD")

    def classify(self, args, ctx):
        try:
            ctx.svc.paths.check_read(args.path)
        except PathDenied as e:
            return Classification(blocked=str(e))
        return Classification()

    def run(self, args, ctx):
        p = ctx.svc.paths.check_read(args.path)
        if not p.is_file():
            return ToolResult.fail(f"{p} is not a file")
        try:
            trades, warns = finance.load_trades(p)
        except ValueError as e:
            return ToolResult.fail(str(e))
        rep = finance.analyze(trades, args.last, args.symbol)
        rep["warnings"] = warns
        if rep["n"] == 0:
            return ToolResult.fail("no trades found" + (f" for {args.symbol}" if args.symbol else ""), report=rep)
        s = rep["stats"]
        pf = "∞" if s["profit_factor_infinite"] else f"{s['profit_factor']:.2f}"
        return ToolResult.success(f"{rep['n']} trades: win rate {s['win_rate']:.1%}, expectancy {s['expectancy']:.2f}, "
                                  f"profit factor {pf}, max drawdown {s['max_drawdown']:.2f}", report=rep, path=str(p))


class FinanceCalc(Tool):
    name = "finance_calc"
    scope = "finance.analysis"
    group = "finance"
    description = "Risk/reward ratio (kind='risk_reward') or position size from account risk (kind='position_size')."

    class Args(BaseModel):
        kind: Literal["risk_reward", "position_size"]
        entry: float
        stop: float
        target: float | None = None
        side: Literal["long", "short"] = "long"
        balance: float | None = None
        risk_percent: float | None = None
        value_per_point: float = 1.0

    def run(self, args, ctx):
        try:
            if args.kind == "risk_reward":
                if args.target is None:
                    return ToolResult.fail("target is required for risk_reward")
                r = finance.risk_reward(args.entry, args.stop, args.target, args.side)
                return ToolResult.success(f"Reward:risk = {r['reward_to_risk']:.2f}:1; break-even win rate {r['breakeven_win_rate']:.1%}", **r)
            if args.balance is None or args.risk_percent is None:
                return ToolResult.fail("balance and risk_percent are required for position_size")
            r = finance.position_size(args.balance, args.risk_percent, args.entry, args.stop, args.value_per_point)
            return ToolResult.success(f"Risking {r['risk_amount']:.2f} with a {r['stop_distance']:g}-point stop → size {r['position_size']:.4f}", **r)
        except ValueError as e:
            return ToolResult.fail(str(e))


class KnowledgeIngest(Tool):
    name = "knowledge_ingest"
    scope = "memory.write"
    group = "memory"
    description = "Add a text/markdown/csv/json/code file to the searchable knowledge base (chunked)."

    class Args(BaseModel):
        path: str
        project: str | None = None

    def classify(self, args, ctx):
        try:
            ctx.svc.paths.check_read(args.path)
        except PathDenied as e:
            return Classification(blocked=str(e))
        return Classification(Risk.LOW, Effect.INTERNAL)

    def run(self, args, ctx):
        p = ctx.svc.paths.check_read(args.path)
        if not p.is_file():
            return ToolResult.fail(f"{p} is not a file")
        if p.suffix.lower() not in ctx.svc.memory.TEXT_SUFFIXES:
            return ToolResult.fail(f"{p.suffix or 'this'} files are not supported yet (text-like files only; PDF/DOCX not implemented)")
        if p.stat().st_size > 2_000_000:
            return ToolResult.fail("file is larger than 2 MB")
        try:
            ids = ctx.svc.memory.ingest_text(p.read_text(encoding="utf-8", errors="replace"), str(p), args.project)
        except PermissionError as e:
            return ToolResult.fail(str(e))
        return ToolResult.success(f"Ingested {p.name} as {len(ids)} chunk(s)", ids=ids, path=str(p))

    def verify(self, args, result, ctx):
        found = [i for i in result.data["ids"] if ctx.svc.memory.get(i)]
        ok = len(found) == len(result.data["ids"]) and bool(found)
        return Verification(ok, "re-read from database", f"{len(found)}/{len(result.data['ids'])} chunks present")


TOOLS = [SpreadsheetSummary, SpreadsheetCalc, FinanceJournal, FinanceCalc, KnowledgeIngest]
